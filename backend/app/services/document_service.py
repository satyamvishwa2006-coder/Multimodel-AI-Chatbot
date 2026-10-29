import os
import io
import csv
import difflib
from typing import Dict, Any, Optional, Tuple, List
from pypdf import PdfReader
from docx import Document as DocxDocument
from app.services.llm_service import get_llm_provider


class DocumentService:
    @staticmethod
    def extract_text(file_path: str, file_type: str) -> str:
        """
        Extracts plain text from PDF, DOCX, TXT, Markdown, CSV, XLSX, and JSON files.
        """
        ext = file_type.lower().lstrip(".")
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        if ext == "pdf":
            reader = PdfReader(file_path)
            extracted = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    extracted.append(text)
            return "\n\n".join(extracted).strip()

        elif ext == "docx":
            doc = DocxDocument(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n\n".join(paragraphs).strip()

        elif ext in ["txt", "md", "json"]:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                return f.read().strip()

        elif ext == "csv":
            lines = []
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                for row in reader:
                    lines.append(", ".join(row))
            return "\n".join(lines).strip()

        elif ext in ["xlsx", "xls"]:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            lines = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                lines.append(f"--- Sheet: {sheet} ---")
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(cell) for cell in row if cell is not None]
                    if row_vals:
                        lines.append(" | ".join(row_vals))
            return "\n".join(lines).strip()

        else:
            raise ValueError(f"Unsupported document format: {ext}")

    @staticmethod
    def edit_content(
        content: str,
        action: str,
        instruction: Optional[str] = None,
        target_language: Optional[str] = "Spanish",
    ) -> str:
        """
        Performs AI document transformation: summarize, simplify, fix_grammar, translate, custom.
        """
        llm = get_llm_provider()

        prompts = {
            "summarize": "Please provide a concise, well-structured summary of the following document while retaining all essential facts and takeaways.",
            "simplify": "Please rewrite the following document into clear, plain, and easy-to-understand language while preserving its core message.",
            "fix_grammar": "Please review and correct all spelling, grammar, punctuation, and flow errors in the following document without changing its underlying facts.",
            "translate": f"Please translate the following text accurately and naturally into {target_language}.",
            "custom": instruction or "Please revise and improve the following text according to best practices.",
        }

        task_prompt = prompts.get(action, prompts["custom"])
        system_instruction = (
            "You are an expert document editor. Modify the user's document according to the requested action. "
            "Return ONLY the revised document text without introductory chatter."
        )

        messages = [
            {"role": "user", "content": f"{task_prompt}\n\n### ORIGINAL DOCUMENT:\n{content}"}
        ]

        revised = llm.generate_response(messages=messages, system_prompt=system_instruction)
        return revised.strip()

    @staticmethod
    def compare_documents(original: str, revised: str) -> Dict[str, Any]:
        """
        Computes a structured diff and change summary between original and revised document versions.
        """
        orig_lines = original.splitlines(keepends=True)
        rev_lines = revised.splitlines(keepends=True)
        diff = list(difflib.unified_diff(orig_lines, rev_lines, fromfile="Original", tofile="Revised"))
        added = sum(1 for line in diff if line.startswith("+") and not line.startswith("+++"))
        removed = sum(1 for line in diff if line.startswith("-") and not line.startswith("---"))

        return {
            "diff_text": "".join(diff),
            "lines_added": added,
            "lines_removed": removed,
            "has_changes": added > 0 or removed > 0,
            "summary": f"{added} lines added, {removed} lines modified/removed.",
        }

    @staticmethod
    def create_content(title: str, prompt: str, file_type: str = "docx") -> str:
        """
        Generates a new structured document from a prompt.
        """
        llm = get_llm_provider()
        system_instruction = (
            "You are a professional document author and data engineer. Create a well-organized, comprehensive document "
            "with appropriate section headings and structured paragraphs (or markdown table data if spreadsheet is requested) "
            "based on the user's prompt."
        )
        messages = [
            {"role": "user", "content": f"Title: {title}\nFile Format: {file_type}\nPrompt: {prompt}"}
        ]
        return llm.generate_response(messages=messages, system_prompt=system_instruction).strip()

    @staticmethod
    def export_document_bytes(content: str, title: str, format_type: str) -> Tuple[bytes, str, str]:
        """
        Exports content into bytes, returning (bytes_data, media_type, filename).
        Supported formats: 'docx', 'pdf', 'xlsx', 'csv', 'md', 'txt'.
        """
        fmt = format_type.lower().lstrip(".")
        clean_title = "".join(c for c in title if c.isalnum() or c in (" ", "_", "-")).strip() or "document"

        if fmt == "docx":
            doc = DocxDocument()
            doc.add_heading(title, level=1)
            for line in content.split("\n\n"):
                clean_line = line.strip()
                if not clean_line:
                    continue
                if clean_line.startswith("# "):
                    doc.add_heading(clean_line[2:], level=1)
                elif clean_line.startswith("## "):
                    doc.add_heading(clean_line[3:], level=2)
                elif clean_line.startswith("### "):
                    doc.add_heading(clean_line[4:], level=3)
                else:
                    doc.add_paragraph(clean_line)

            buffer = io.BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            return (
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                f"{clean_title}.docx",
            )

        elif fmt == "pdf":
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib import colors

            buffer = io.BytesIO()
            doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
            styles = getSampleStyleSheet()

            title_style = ParagraphStyle(
                "DocTitle",
                parent=styles["Title"],
                fontName="Helvetica-Bold",
                fontSize=20,
                leading=24,
                textColor=colors.HexColor("#1e293b"),
                spaceAfter=14,
            )
            body_style = ParagraphStyle(
                "DocBody",
                parent=styles["Normal"],
                fontName="Helvetica",
                fontSize=10,
                leading=14,
                textColor=colors.HexColor("#334155"),
                spaceAfter=8,
            )

            story = [Paragraph(title, title_style), Spacer(1, 12)]
            for para in content.split("\n\n"):
                p_clean = para.strip().replace("\n", "<br/>")
                if p_clean:
                    story.append(Paragraph(p_clean, body_style))
                    story.append(Spacer(1, 6))

            doc.build(story)
            buffer.seek(0)
            return (
                buffer.getvalue(),
                "application/pdf",
                f"{clean_title}.pdf",
            )

        elif fmt in ["xlsx", "xls"]:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = clean_title[:30]

            ws.append([title])
            ws.append([])

            for line in content.split("\n"):
                clean = line.strip()
                if not clean:
                    continue
                # If pipe-separated markdown table row
                if "|" in clean:
                    cells = [c.strip() for c in clean.split("|") if c.strip() and not set(c.strip()).issubset({"-", ":"})]
                    if cells:
                        ws.append(cells)
                # If comma-separated
                elif "," in clean:
                    cells = [c.strip() for c in clean.split(",")]
                    ws.append(cells)
                else:
                    ws.append([clean])

            buffer = io.BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            return (
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                f"{clean_title}.xlsx",
            )

        elif fmt == "csv":
            buffer = io.StringIO()
            writer = csv.writer(buffer)
            writer.writerow([title])
            writer.writerow([])
            for line in content.split("\n"):
                clean = line.strip()
                if not clean:
                    continue
                if "|" in clean:
                    cells = [c.strip() for c in clean.split("|") if c.strip() and not set(c.strip()).issubset({"-", ":"})]
                    if cells:
                        writer.writerow(cells)
                elif "," in clean:
                    cells = [c.strip() for c in clean.split(",")]
                    writer.writerow(cells)
                else:
                    writer.writerow([clean])

            return (
                buffer.getvalue().encode("utf-8"),
                "text/csv; charset=utf-8",
                f"{clean_title}.csv",
            )

        elif fmt == "md":
            data = f"# {title}\n\n{content}".encode("utf-8")
            return (
                data,
                "text/markdown; charset=utf-8",
                f"{clean_title}.md",
            )

        else:  # 'txt'
            data = f"{title}\n{'=' * len(title)}\n\n{content}".encode("utf-8")
            return (
                data,
                "text/plain; charset=utf-8",
                f"{clean_title}.txt",
            )
